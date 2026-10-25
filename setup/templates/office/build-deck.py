#!/usr/bin/env python3
"""Slide deck builder — HTML slide -> PNG (headless Chrome) -> PPTX.

Why this shape: python-pptx has no layout engine, so hand-assembled native
shapes top out at draft quality. Authoring each slide as a 1920x1080 HTML
document and dropping the render in full-bleed gives real typography and
layout, at the cost of non-editable slides. Detail goes in the speaker notes
and in the companion document — never on the slide face.

Design law (see ~/.claude/knowledge/presentation-design.md):
  <= 7 words on a slide, big type, generous air, visual not textual.

Use as a library:

    from build_deck import Deck
    d = Deck("My Topic")
    d.cover("AGENT FLEET", "One Big Idea", "subtitle line", "footer line")
    d.statement("KICKER", "Seven words or fewer here", notes="...")
    d.stats("KICKER", "The numbers", [("18", "built", ""), ("7", "left", "ok")])
    d.build("out.pptx")

Or run it directly for a self-test deck:

    python3 build-deck.py [out.pptx]
"""
import html
import os
import shutil
import subprocess
import sys
import tempfile

from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import tokens as T  # noqa: E402

_f = html.escape


# ------------------------------------------------------------- environment --
def is_wsl():
    try:
        with open("/proc/version") as fh:
            return "microsoft" in fh.read().lower()
    except OSError:
        return False


def find_chrome():
    """Return (executable, needs_windows_paths)."""
    if is_wsl():
        for p in ("/mnt/c/Program Files/Google/Chrome/Application/chrome.exe",
                  "/mnt/c/Program Files (x86)/Google/Chrome/Application/chrome.exe",
                  "/mnt/c/Program Files/Microsoft/Edge/Application/msedge.exe"):
            if os.path.exists(p):
                return p, True
    for name in ("google-chrome", "chromium", "chromium-browser", "chrome"):
        p = shutil.which(name)
        if p:
            return p, False
    raise RuntimeError(
        "no Chrome/Chromium found. Install chromium, or on WSL install "
        "Google Chrome on the Windows side.")


CSS = f"""
*{{margin:0;padding:0;box-sizing:border-box}}
html,body{{width:{T.CANVAS_W}px;height:{T.CANVAS_H}px;overflow:hidden}}
body{{font-family:{T.FONT_STACK};background:#{T.SURFACE};color:#{T.INK};
     -webkit-font-smoothing:antialiased}}
.slide{{width:{T.CANVAS_W}px;height:{T.CANVAS_H}px;position:relative;display:flex;
       flex-direction:column;justify-content:center}}
.pad{{padding:{T.PAD_V}px {T.PAD_H}px}}
.kicker{{font-size:26px;font-weight:700;letter-spacing:3px;text-transform:uppercase;
        color:#{T.PRIMARY};margin-bottom:26px}}
.accent{{width:120px;height:8px;border-radius:4px;background:#{T.PRIMARY};margin:0 0 44px 0}}
.h1{{font-size:88px;font-weight:800;line-height:1.05;letter-spacing:-1px;color:#{T.PRIMARY_DEEP}}}
.sub{{font-size:38px;font-weight:400;color:#{T.MUTED};margin-top:34px}}
.footer{{position:absolute;left:{T.PAD_H}px;bottom:64px;font-size:20px;color:#{T.FAINT_INK}}}
.src{{position:absolute;right:60px;bottom:34px;font-size:17px;font-style:italic;color:#{T.FAINT_INK}}}
.cover{{background:linear-gradient(125deg,#{T.PRIMARY} 0%,#{T.PRIMARY_DEEP} 100%);color:#fff}}
.cover .brand{{position:absolute;top:90px;left:{T.PAD_H}px;font-size:26px;font-weight:700;
              letter-spacing:4px;text-transform:uppercase;color:#{T.SURFACE_TINT}}}
.cover .accent{{background:#{T.ACCENT}}}
.cover .big{{font-size:130px;font-weight:800;letter-spacing:-2px;line-height:1}}
.cover .tag{{font-size:46px;font-weight:300;color:#{T.SURFACE_TINT};margin-top:30px}}
.cover .foot{{position:absolute;left:{T.PAD_H}px;bottom:90px;font-size:24px;color:#{T.PRIMARY_SOFT}}}
.row{{display:flex;gap:56px;margin-top:70px}}
.stat{{flex:1;min-width:340px;border-radius:14px;padding:48px 44px;
      background:#{T.SURFACE_TINT};border-left:10px solid #{T.PRIMARY}}}
.stat.ok{{border-left-color:#{T.ACCENT}}}
.stat .num{{font-size:92px;font-weight:800;letter-spacing:-2px;line-height:1;color:#{T.PRIMARY}}}
.stat.ok .num{{color:#{T.ACCENT}}}
.stat .lab{{font-size:30px;font-weight:500;color:#{T.MUTED};margin-top:20px}}
.chip{{flex:1;min-width:360px;border-radius:16px;padding:48px 44px;
      background:#{T.SURFACE_TINT};border:2px solid #{T.BORDER}}}
.chip.muted{{background:#{T.SURFACE_MUTE};border-color:#{T.BORDER_MUTE}}}
.chip .ct{{font-size:40px;font-weight:800;color:#{T.PRIMARY_DEEP}}}
.chip.muted .ct{{color:#{T.FAINT_INK}}}
.chip .cs{{font-size:26px;line-height:1.35;margin-top:16px;color:#{T.MUTED}}}
.cols{{display:flex;gap:70px;margin-top:70px}}
.col{{flex:1}}
.col h3{{font-size:34px;font-weight:800;margin-bottom:30px;color:#{T.ACCENT}}}
.col.later h3{{color:#{T.FAINT_INK}}}
.col li{{list-style:none;font-size:30px;margin:18px 0;padding-left:44px;position:relative;
        color:#{T.INK}}}
.col li:before{{content:"";position:absolute;left:0;top:14px;width:18px;height:18px;
               border-radius:5px;background:#{T.ACCENT}}}
.col.later li:before{{background:#{T.BORDER}}}
.imghead{{padding:70px 90px 40px 90px}}
.imghead .t{{font-size:60px;font-weight:800;letter-spacing:-1px;color:#{T.PRIMARY_DEEP}}}
.imgwrap{{flex:1;display:flex;align-items:center;justify-content:center;padding:0 90px 70px 90px}}
.imgwrap img{{max-width:100%;max-height:100%;border-radius:12px;border:1px solid #{T.BORDER};
             box-shadow:0 24px 60px rgba(10,61,70,.16)}}
.imgwrap.diag img{{border:none;border-radius:0;box-shadow:none}}
.cap{{position:absolute;right:90px;bottom:30px;font-size:19px;font-style:italic;color:#{T.FAINT_INK}}}
.closing li{{list-style:none;font-size:44px;font-weight:300;margin:22px 0;color:#{T.SURFACE_TINT}}}
.closing .big{{font-size:96px;font-weight:800}}
"""


def page(body_html, extra_class=""):
    return (f"<!doctype html><html><head><meta charset='utf-8'>"
            f"<style>{CSS}</style></head><body>"
            f"<div class='slide {extra_class}'>{body_html}</div></body></html>")


# ------------------------------------------------------------------- deck ---
class Deck:
    """Collects (html, speaker_notes) pairs, then renders and assembles."""

    def __init__(self, name, assets=None):
        self.name = name
        self.slides = []
        self.assets = assets or []      # extra image paths to make available

    def _add(self, html_body, notes, cls=""):
        self.slides.append((page(html_body, cls), notes or ""))
        return self

    # -- slide types --------------------------------------------------------
    def cover(self, brand, big, tag="", foot="", notes=""):
        b = (f"<div class='pad'><div class='brand'>{_f(brand)}</div>"
             f"<div class='accent'></div><div class='big'>{_f(big)}</div>")
        if tag:
            b += f"<div class='tag'>{_f(tag)}</div>"
        b += "</div>"
        if foot:
            b += f"<div class='foot'>{_f(foot)}</div>"
        return self._add(b, notes, "cover")

    def statement(self, kicker, h1, sub="", foot="", src="", notes=""):
        b = "<div class='pad'>"
        if kicker:
            b += f"<div class='kicker'>{_f(kicker)}</div>"
        b += f"<div class='accent'></div><div class='h1'>{_f(h1)}</div>"
        if sub:
            b += f"<div class='sub'>{_f(sub)}</div>"
        b += "</div>"
        if foot:
            b += f"<div class='footer'>{_f(foot)}</div>"
        if src:
            b += f"<div class='src'>{_f(src)}</div>"
        return self._add(b, notes)

    def stats(self, kicker, h1, items, src="", notes=""):
        """items: [(number, label, cssclass)] — cssclass '' or 'ok'."""
        cards = "".join(
            f"<div class='stat {c}'><div class='num'>{_f(n)}</div>"
            f"<div class='lab'>{_f(l)}</div></div>" for n, l, c in items)
        b = (f"<div class='pad'><div class='kicker'>{_f(kicker)}</div>"
             f"<div class='h1'>{_f(h1)}</div><div class='row'>{cards}</div></div>")
        if src:
            b += f"<div class='src'>{_f(src)}</div>"
        return self._add(b, notes)

    def chips(self, kicker, h1, items, notes=""):
        """items: [(title, detail, cssclass)] — cssclass '' or 'muted'."""
        cs = "".join(
            f"<div class='chip {c}'><div class='ct'>{_f(t)}</div>"
            f"<div class='cs'>{_f(d)}</div></div>" for t, d, c in items)
        b = (f"<div class='pad'><div class='kicker'>{_f(kicker)}</div>"
             f"<div class='h1'>{_f(h1)}</div><div class='row'>{cs}</div></div>")
        return self._add(b, notes)

    def twocol(self, h1, now_title, now, later_title, later, notes=""):
        def ul(items):
            return "".join(f"<li>{_f(i)}</li>" for i in items)
        b = (f"<div class='pad'><div class='h1'>{_f(h1)}</div><div class='cols'>"
             f"<div class='col'><h3>{_f(now_title)}</h3><ul>{ul(now)}</ul></div>"
             f"<div class='col later'><h3>{_f(later_title)}</h3><ul>{ul(later)}</ul></div>"
             f"</div></div>")
        return self._add(b, notes)

    def image(self, kicker, title, img_path, cap="", diag=False, notes=""):
        self.assets.append(img_path)
        base = os.path.basename(img_path)
        b = (f"<div class='imghead'><div class='kicker'>{_f(kicker)}</div>"
             f"<div class='t'>{_f(title)}</div></div>"
             f"<div class='imgwrap {'diag' if diag else ''}'>"
             f"<img src='assets/{_f(base)}'></div>")
        if cap:
            b += f"<div class='cap'>{_f(cap)}</div>"
        return self._add(b, notes)

    def closing(self, h1, items, notes=""):
        li = "".join(f"<li>&#8594;&nbsp;&nbsp;{_f(i)}</li>" for i in items)
        b = (f"<div class='pad'><div class='accent'></div>"
             f"<div class='big'>{_f(h1)}</div><ul class='closing'>{li}</ul></div>")
        return self._add(b, notes, "cover closing")

    # -- render -------------------------------------------------------------
    def build(self, out_path, keep_work=False):
        if not self.slides:
            raise ValueError("deck has no slides")
        chrome, win_paths = find_chrome()

        if win_paths:
            # Windows Chrome cannot read WSL /home paths — work under /mnt/c.
            work = "/mnt/c/temp/fleet-render"
            url_base = "file:///C:/temp/fleet-render/"
        else:
            work = tempfile.mkdtemp(prefix="fleet-render-")
            url_base = "file://" + work + "/"
        shutil.rmtree(work, ignore_errors=True)
        os.makedirs(os.path.join(work, "assets"), exist_ok=True)
        for a in self.assets:
            if os.path.exists(a):
                shutil.copy2(a, os.path.join(work, "assets", os.path.basename(a)))

        pngs = []
        for idx, (html_doc, notes) in enumerate(self.slides):
            stem = f"slide_{idx:02d}"
            with open(os.path.join(work, stem + ".html"), "w",
                      encoding="utf-8") as fh:
                fh.write(html_doc)
            png = os.path.join(work, stem + ".png")
            shot = (f"C:\\temp\\fleet-render\\{stem}.png" if win_paths else png)
            proc = subprocess.run(
                [chrome, "--headless=new", "--disable-gpu", "--hide-scrollbars",
                 "--force-device-scale-factor=1", f"--screenshot={shot}",
                 f"--window-size={T.CANVAS_W},{T.CANVAS_H}",
                 url_base + stem + ".html"],
                capture_output=True, text=True)
            # The original of this pipeline swallowed Chrome's exit code and
            # inserted a blank slide on failure while still reporting success.
            if proc.returncode != 0:
                raise RuntimeError(
                    f"chrome failed on {stem} (rc={proc.returncode}): "
                    f"{proc.stderr.strip()[:400]}")
            if not os.path.exists(png) or os.path.getsize(png) < 2000:
                raise RuntimeError(f"chrome produced no usable PNG for {stem}")
            pngs.append((png, notes))

        prs = Presentation()
        prs.slide_width = Inches(T.SLIDE_W_IN)
        prs.slide_height = Inches(T.SLIDE_H_IN)
        blank = prs.slide_layouts[6]
        for png, notes in pngs:
            s = prs.slides.add_slide(blank)
            s.shapes.add_picture(png, 0, 0, width=prs.slide_width,
                                 height=prs.slide_height)
            s.notes_slide.notes_text_frame.text = notes
        prs.save(out_path)

        if not keep_work:
            shutil.rmtree(work, ignore_errors=True)
        return len(pngs)


def _self_test(out):
    d = Deck("self-test")
    d.cover("AGENT FLEET", "Office Template", "Generic house style",
            "Replace the logo during onboarding", notes="Cover.")
    d.statement("PRINCIPLE", "Seven words is the limit",
                sub="Detail belongs in the talk and the handout.",
                foot="Agent Fleet", src="presentation-design.md")
    d.stats("COVERAGE", "Where the work stands",
            [("18", "built", ""), ("5", "partial", ""), ("7", "unbuilt", "ok")])
    d.chips("SCOPE", "Two candidate scopes",
            [("Combined core", "Shows the concept. Carries nothing through a gate.", ""),
             ("Pilot-capable", "The smallest scope a real project can run on.", "muted")])
    d.twocol("What is in, what is later", "Now",
             ["Portfolio board", "Dependency register"], "Later",
             ["Resource simulation", "Novice mode"])
    d.closing("Decisions requested", ["Confirm the scope", "Nominate the pilot",
                                      "Approve the baseline study"])
    n = d.build(out)
    print(f"built: {out}  {n} slides")


if __name__ == "__main__":
    _self_test(sys.argv[1] if len(sys.argv) > 1 else "fleet-deck-selftest.pptx")
