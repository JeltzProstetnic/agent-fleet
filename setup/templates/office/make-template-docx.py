#!/usr/bin/env python3
"""Generate `assets/fleet-template.docx` — the branded reference document.

`build-docx.py` opens this file, clears its body and writes new content into it,
which is how the styles, the header logo and the footer page-number fields
survive into every deliverable.

Run once after changing `tokens.py` or replacing `assets/logo.svg`:

    python3 make-template-docx.py

Idempotent: it overwrites the template from scratch every time.
"""
import os
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import tokens as T  # noqa: E402

OUT = os.path.join(HERE, "assets", "fleet-template.docx")
TITLE_PLACEHOLDER = "{{DOCUMENT_TITLE}}"


def rgb(h):
    return RGBColor.from_string(h)


def _field(paragraph, instr):
    """Insert a Word field (PAGE, NUMPAGES) that recalculates on open."""
    run = paragraph.add_run()
    for tag, attr, text in (("w:fldChar", "w:fldCharType", "begin"),
                            ("w:instrText", None, f" {instr} "),
                            ("w:fldChar", "w:fldCharType", "end")):
        el = OxmlElement(tag)
        if attr:
            el.set(qn(attr), text)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _style(doc, name, *, font=None, size=None, bold=None, color=None,
           space_before=None, space_after=None, keep_next=None):
    st = doc.styles[name]
    f = st.font
    if font:
        f.name = font
        # East-Asian + complex-script names must be set too or Word substitutes
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(a), font)
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        f.color.rgb = rgb(color)
    pf = st.paragraph_format
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if space_after is not None:
        pf.space_after = Pt(space_after)
    if keep_next is not None:
        pf.keep_with_next = keep_next
    return st


def build():
    doc = Document()

    # ---- page geometry -----------------------------------------------------
    sec = doc.sections[0]
    sec.page_width = Mm(T.PAGE_W_MM)
    sec.page_height = Mm(T.PAGE_H_MM)
    sec.top_margin = Mm(T.MARGIN_TOP_MM)
    sec.bottom_margin = Mm(T.MARGIN_BOT_MM)
    sec.left_margin = Mm(T.MARGIN_SIDE_MM)
    sec.right_margin = Mm(T.MARGIN_SIDE_MM)
    sec.header_distance = Mm(T.HEADER_MM)
    sec.footer_distance = Mm(T.FOOTER_MM)

    # ---- styles ------------------------------------------------------------
    _style(doc, "Normal", font=T.FONT_BODY, size=T.PT_BODY, color=T.INK)
    _style(doc, "Title", font=T.FONT_HEAD, size=T.PT_TITLE, bold=True,
           color=T.PRIMARY, space_after=12)
    _style(doc, "Heading 1", font=T.FONT_HEAD, size=T.PT_H1, bold=True,
           color=T.PRIMARY, space_before=18, space_after=6, keep_next=True)
    _style(doc, "Heading 2", font=T.FONT_HEAD, size=T.PT_H2, bold=True,
           color=T.PRIMARY, space_before=14, space_after=4, keep_next=True)
    _style(doc, "Heading 3", font=T.FONT_HEAD, size=T.PT_H3, bold=True,
           color=T.PRIMARY_DEEP, space_before=10, space_after=3, keep_next=True)
    _style(doc, "List Paragraph", font=T.FONT_BODY, size=T.PT_BODY)
    for nm in ("Header", "Footer"):
        _style(doc, nm, font=T.FONT_BODY, size=T.PT_SMALL, color=T.MUTED)
    # Quote style used for blockquotes / draft warnings
    _style(doc, "Quote", font=T.FONT_BODY, size=T.PT_BODY, color=T.PRIMARY_DEEP)

    text_w = Mm(T.TEXT_W_MM)

    # ---- header: logo flush right, then the document-title line ------------
    hdr = sec.header
    hdr.is_linked_to_previous = False
    p_logo = hdr.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_logo.paragraph_format.space_after = Pt(2)
    logo = os.path.join(HERE, T.LOGO_PNG)
    if os.path.exists(logo):
        p_logo.add_run().add_picture(logo, width=Mm(T.LOGO_W_MM),
                                     height=Mm(T.LOGO_H_MM))
    p_title = hdr.add_paragraph(style="Header")
    p_title.paragraph_format.tab_stops.add_tab_stop(text_w,
                                                    WD_TAB_ALIGNMENT.RIGHT)
    p_title.add_run(f"Document title: {TITLE_PLACEHOLDER}")
    _rule(p_title)

    # ---- footer: status | version | page N of M ----------------------------
    ftr = sec.footer
    ftr.is_linked_to_previous = False
    p_f = ftr.paragraphs[0]
    p_f.style = doc.styles["Footer"]
    ts = p_f.paragraph_format.tab_stops
    ts.add_tab_stop(Mm(T.TEXT_W_MM / 2), WD_TAB_ALIGNMENT.CENTER)
    ts.add_tab_stop(text_w, WD_TAB_ALIGNMENT.RIGHT)
    p_f.add_run("{{STATUS}}\t{{VERSION}}\tpage ")
    _field(p_f, "PAGE \\* Arabic")
    p_f.add_run(" of ")
    _field(p_f, "NUMPAGES \\* Arabic")

    # ---- a single body paragraph; build-docx.py clears it ------------------
    doc.add_paragraph("")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    return OUT


def _rule(paragraph):
    """Thin bottom border under a paragraph — the header rule."""
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")        # 0.5 pt
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), T.BORDER)
    borders.append(bottom)
    ppr.append(borders)


if __name__ == "__main__":
    print("built:", build())
